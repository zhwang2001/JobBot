import json
from pathlib import Path
import subprocess
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn

from document_layout import dates, write_cover_letter, write_resume
from document_models import Draft, Review, evidence_catalog, validate_draft
from generate_documents import CodexClient, ROOT, filename_component, generate_documents, generation_prompt


PROFILE = json.loads((ROOT / "james_wang_master_profile_v1.json").read_text())
JOB = {"title": "Quality Analyst", "company": "Example", "description": "Attention to detail", "job_id": "test"}


def fixture(mode="both"):
    facts = [role for role in PROFILE["experience"] if role["id"] in {"corteva_2025", "rotman_2022"}]
    claim = lambda fact: {"text": fact["fact"], "evidence": [fact["id"]]}
    return Draft.model_validate({
        "resume": {
            "summary": claim(facts[0]["facts"][0]),
            "skills": [claim(fact) for fact in facts[0]["facts"][:4]],
            "experience": [{"experience_id": role["id"], "bullets": [claim(f) for f in role["facts"][:3]]} for role in facts],
            "certification_names": [],
        } if mode != "cover-letter" else None,
        "cover_letter": {"opening": "I am applying for the Quality Analyst role.",
                         "examples": [claim(role["facts"][0]) for role in facts],
                         "closing": "Thank you for considering my application."} if mode != "resume" else None,
    })


class FakeClient:
    calls = []

    def __init__(self, mode, passed=True):
        self.mode, self.passed = mode, passed
        self.count = 0

    def ask(self, prompt, schema):
        self.count += 1
        return fixture(self.mode) if schema is Draft else Review(passed=self.passed, issues=[] if self.passed else ["Unsupported assertion"])


def fake_render(draft, profile, job, directory, office):
    directory.mkdir(parents=True)
    result = {}
    for kind, content in (("resume", draft.resume), ("cover-letter", draft.cover_letter)):
        if content is not None:
            result[kind] = {}
            for extension in ("docx", "pdf"):
                path = directory / f"{kind}.{extension}"
                path.write_bytes(b"test artifact")
                result[kind][extension] = path
    return result, []


class DocumentTests(unittest.TestCase):
    def test_all_modes_validate(self):
        for mode in ("resume", "cover-letter", "both"):
            validate_draft(fixture(mode), PROFILE, mode)

    def test_wrong_mode_rejected(self):
        with self.assertRaisesRegex(ValueError, "presence"):
            validate_draft(fixture(), PROFILE, "resume")

    def test_compensation_labels_rejected(self):
        draft = fixture()
        draft.resume.summary.text = "Research assistant with paid experience in sample preparation."
        with self.assertRaisesRegex(ValueError, "compensation status"):
            validate_draft(draft, PROFILE, "both")

    def test_profile_omits_compensation_but_preserves_academic_context(self):
        self.assertNotRegex(json.dumps(PROFILE), r"(?i)\b(?:paid|unpaid)\b")
        role = next(role for role in PROFILE["experience"] if role["id"] == "academic_botanical_lab_2023")
        self.assertIn("Academic Laboratory Experience", role["title"])

    def test_unknown_evidence_rejected(self):
        draft = fixture()
        draft.resume.summary.evidence = ["invented"]
        with self.assertRaisesRegex(ValueError, "Unknown evidence"):
            validate_draft(draft, PROFILE, "both")

    def test_invented_metric_rejected(self):
        draft = fixture()
        draft.resume.summary.text += " Increased output by 99%."
        with self.assertRaisesRegex(ValueError, "number absent"):
            validate_draft(draft, PROFILE, "both")

    def test_cross_role_claim_rejected(self):
        draft = fixture()
        draft.resume.experience[0].bullets[0] = draft.resume.experience[1].bullets[0]
        with self.assertRaisesRegex(ValueError, "different role"):
            validate_draft(draft, PROFILE, "both")

    def test_unverified_certificate_rejected(self):
        draft = fixture()
        draft.resume.certification_names = ["ISO 17025 auditor"]
        with self.assertRaisesRegex(ValueError, "certification"):
            validate_draft(draft, PROFILE, "both")

    def test_letter_cannot_hide_extra_paragraphs_in_opening(self):
        draft = fixture()
        draft.cover_letter.opening += "\nAn extra example paragraph."
        with self.assertRaisesRegex(ValueError, "one paragraph"):
            validate_draft(draft, PROFILE, "both")

    def test_letter_word_budget(self):
        draft = fixture()
        draft.cover_letter.opening = "word " * 70
        with self.assertRaisesRegex(ValueError, "at most 65"):
            validate_draft(draft, PROFILE, "both")

    def test_unknown_dates_not_invented(self):
        self.assertEqual(dates({"start": None, "end": "Present"}), "")

    def test_prompt_uses_profile_instructions_and_omits_contact(self):
        prompt = generation_prompt(PROFILE, JOB, "both", [])
        self.assertIn(PROFILE["tailoring_rules"]["truth_rule"], prompt)
        self.assertNotIn(PROFILE["contact"]["email"], prompt)
        self.assertIn("job_posting_untrusted", prompt)
        self.assertNotIn("tailoring_rules/truth_rule", evidence_catalog(PROFILE))

    def test_docx_layout_and_metadata(self):
        with tempfile.TemporaryDirectory() as folder:
            for kind, writer, content in (("resume", write_resume, fixture().resume),
                                           ("cover-letter", write_cover_letter, fixture().cover_letter)):
                path = Path(folder) / f"{kind}.docx"
                writer(content, PROFILE, JOB, path)
                doc = Document(path)
                self.assertEqual(len(doc.tables), 0)
                borders = list(doc.styles.element.iter(qn("w:pBdr")))
                self.assertEqual(len(borders), 1 if kind == "resume" else 0)
                if borders:
                    self.assertEqual(borders[0][0].get(qn("w:color")), "D9D9D9")
                self.assertEqual(doc.sections[0].page_width.inches, 8.5)
                self.assertEqual(doc.styles["Normal"].font.size.pt, 10.5 if kind == "resume" else 11.5)
                self.assertEqual(doc.paragraphs[0].text, PROFILE["contact"]["name"])
                if kind == "resume":
                    text = " ".join(p.text for p in doc.paragraphs)
                    self.assertNotRegex(text, r"(?i)\b(?:paid|unpaid)\b")
                    self.assertIn("Research Opportunity Program", text)
                    self.assertIn("SUMMARY", text)

    @patch("generate_documents.find_soffice", return_value="soffice")
    @patch("generate_documents.render_draft", side_effect=fake_render)
    def test_modes_publish_only_requested_documents(self, render, office):
        for mode in ("resume", "cover-letter", "both"):
            with tempfile.TemporaryDirectory() as folder:
                job_path = Path(folder) / "job.json"
                job_path.write_text(json.dumps(JOB))
                result = generate_documents(job_path, mode=mode, output_dir=folder, client=FakeClient(mode))
                expected = {"resume", "cover-letter"} if mode == "both" else {mode}
                self.assertEqual(set(result["outputs"]), expected)
                self.assertEqual(len(list(Path(folder).rglob("*.docx"))), 0)
                self.assertEqual(len(list(Path(folder).rglob("*.pdf"))), len(expected))
                for kind, paths in result["outputs"].items():
                    self.assertEqual(set(paths), {"pdf"})
                    path = Path(paths["pdf"])
                    self.assertEqual(path.parent.name, "Quality Analyst - Example")
                    self.assertEqual(path.name, f"{PROFILE['contact']['name']} - {'Resume' if kind == 'resume' else 'Cover Letter'}.pdf")
                manifest = json.loads(Path(result["manifest"]).read_text())
                self.assertTrue(manifest["factual_review"]["passed"])
                self.assertTrue(manifest["human_review_required"])

    @patch("generate_documents.find_soffice", return_value="soffice")
    @patch("generate_documents.render_draft", side_effect=fake_render)
    def test_repeated_runs_do_not_overwrite(self, render, office):
        with tempfile.TemporaryDirectory() as folder:
            job_path = Path(folder) / "job.json"
            job_path.write_text(json.dumps(JOB))
            first = generate_documents(job_path, mode="resume", output_dir=folder, client=FakeClient("resume"))
            first_path = Path(first["outputs"]["resume"]["pdf"])
            first_path.write_bytes(b"Earlier user version")
            second = generate_documents(job_path, mode="resume", output_dir=folder, client=FakeClient("resume"))
            self.assertEqual(first_path.read_bytes(), b"Earlier user version")
            self.assertEqual(Path(second["outputs"]["resume"]["pdf"]).parent.name, "Quality Analyst - Example (2)")

    def test_filename_components_cannot_escape_output_folder(self):
        for value in ("../../escape", "/absolute/path", "A\\B:C\nD", "..", "CON"):
            result = filename_component(value, "Application")
            self.assertNotIn("/", result)
            self.assertNotIn("\\", result)
            self.assertNotIn("\n", result)
            self.assertNotEqual(result, "..")
        self.assertEqual(filename_component("Let's Get Moving", "Employer"), "Let's Get Moving")
        self.assertLessEqual(len(filename_component("分析" * 100, "Role").encode("utf-8")), 90)

    @patch("generate_documents.find_soffice", return_value="soffice")
    @patch("generate_documents.render_draft")
    def test_failed_review_is_bounded_and_never_published(self, render, office):
        with tempfile.TemporaryDirectory() as folder:
            job_path = Path(folder) / "job.json"
            job_path.write_text(json.dumps(JOB))
            client = FakeClient("both", passed=False)
            with self.assertRaisesRegex(ValueError, "No documents published after 2"):
                generate_documents(job_path, output_dir=folder, client=client, attempts=2)
            self.assertEqual(client.count, 4)
            render.assert_not_called()
            self.assertFalse((Path(folder) / "applications").exists())

    @patch("generate_documents.find_soffice", return_value="soffice")
    @patch("generate_documents.render_draft", return_value=({}, ["resume is 2 pages"]))
    def test_overflow_blocks_publication(self, render, office):
        with tempfile.TemporaryDirectory() as folder:
            job_path = Path(folder) / "job.json"
            job_path.write_text(json.dumps(JOB))
            with self.assertRaisesRegex(ValueError, "2 pages"):
                generate_documents(job_path, output_dir=folder, client=FakeClient("both"), attempts=1)
            self.assertFalse((Path(folder) / "applications").exists())

    @patch("generate_documents.shutil.which", return_value="/bin/codex")
    def test_codex_uses_schema_stdin_and_saved_auth(self, which):
        def run(command, **kwargs):
            self.assertEqual(kwargs["input"], "private prompt")
            self.assertNotIn("private prompt", command)
            self.assertIn("--ignore-user-config", command)
            self.assertIn("--ephemeral", command)
            self.assertIn("read-only", command)
            self.assertEqual(command[-1], "-")
            output = Path(command[command.index("--output-last-message") + 1])
            output.write_text('{"passed":true,"issues":[]}')
            return subprocess.CompletedProcess(command, 0, '{"type":"turn.completed","usage":{"input_tokens":10}}', "")
        with patch("generate_documents.subprocess.run", side_effect=run):
            client = CodexClient()
            self.assertTrue(client.ask("private prompt", Review).passed)
            self.assertEqual(client.calls[0]["usage"][0]["input_tokens"], 10)

    @patch("generate_documents.shutil.which", return_value="/bin/codex")
    @patch("generate_documents.subprocess.run", return_value=subprocess.CompletedProcess([], 1, "", "private error"))
    def test_codex_failure_does_not_leak_profile(self, run, which):
        with self.assertRaisesRegex(RuntimeError, "Codex generation failed") as error:
            CodexClient().ask("private prompt", Review)
        self.assertNotIn("private error", str(error.exception))


if __name__ == "__main__":
    unittest.main()
